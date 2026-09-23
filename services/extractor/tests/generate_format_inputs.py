"""Synthetic inputs for scripts/formats-test.sh; never application demo answers."""
import base64
import io
import json
import sys

import docx
import openpyxl

sys.path.insert(0, "/app")
from test_extract import _pdf


def make(kind, lines):
    if kind == "pdf":
        return _pdf([f"1.{i}. {line}" for i, line in enumerate(lines, 1)])
    out = io.BytesIO()
    if kind == "xlsx":
        book = openpyxl.Workbook()
        book.active.title = "Обязанности"
        for line in lines:
            book.active.append([line])
        book.save(out)
    else:
        document = docx.Document()
        if kind == "docx-table":
            table = document.add_table(rows=len(lines), cols=1)
            for row, line in zip(table.rows, lines):
                row.cells[0].text = line
        else:
            for line in lines:
                document.add_paragraph(line)
        document.save(out)
    return out.getvalue()


before = ["Audit department (AUD) conducts annual audits.", "Audit department (AUD) prepares quarterly reports."]
after = ["Audit department (AUD) conducts annual audits.", "Risk department (RISK) maintains the risk register."]
cases = []
for kind in ["docx", "docx-table", "pdf", "xlsx"]:
    ext = "docx" if kind.startswith("docx") else kind
    files = [{"side": side, "filename": f"{side}.{ext}", "data": base64.b64encode(make(kind, lines)).decode()}
             for side, lines in [("before", before), ("after", after)]]
    cases.append({"kind": kind, "files": files})
print(json.dumps(cases))
