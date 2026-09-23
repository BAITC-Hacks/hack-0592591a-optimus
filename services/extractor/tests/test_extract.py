import io
import zipfile

import docx
import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.main import app
from app.model import clause_from_text

client = TestClient(app)


def _post(name: str, data: bytes):
    return client.post("/extract", files={"file": (name, data, "application/octet-stream")})


def _docx() -> bytes:
    document = docx.Document()
    document.add_heading("Положение об отделе закупок", level=1)
    document.add_heading("Функции", level=2)
    document.add_paragraph("Планирование закупок.", style="List Number")
    document.add_paragraph("Проведение тендеров.", style="List Number")
    document.add_paragraph("3.5. Контроль исполнения договоров.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Подразделение"
    table.cell(0, 1).text = "Функция"
    table.cell(1, 0).text = "Отдел закупок"
    table.cell(1, 1).text = "Закупки"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _xlsx() -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Структура"
    sheet.append(["Код", "Подразделение", "Подчинение"])
    sheet.append([None, None, None])
    sheet.append([101, "Отдел закупок", "Финансовый департамент"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _pdf(lines: list[str]) -> bytes:
    """Minimal one-page PDF with a Helvetica text layer (ASCII only)."""
    escaped = [line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)") for line in lines]
    content = "BT /F1 12 Tf 72 720 Td 14 TL " + " ".join(f"({line}) Tj T*" for line in escaped) + " ET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        f"<< /Length {len(content)} >>\nstream\n{content}\nendstream",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{body}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += "".join(f"{offset:010d} 00000 n \n" for offset in offsets).encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return out


def test_health():
    assert client.get("/health").json() == {"status": "ok"}


@pytest.mark.parametrize(
    ("text", "clause"),
    [
        ("3.2. Функции отдела", "3.2"),
        ("3.2.1 Подпункт", "3.2.1"),
        ("4) Иное", "4"),
        ("Глава 2. Задачи", "Глава 2"),
        ("15 сотрудников в штате", None),
        ("2026 год", None),
    ],
)
def test_clause_from_text(text, clause):
    assert clause_from_text(text) == clause


def test_docx_keeps_order_numbering_sections_and_tables():
    body = _post("polozhenie.docx", _docx()).json()
    assert body["document"]["format"] == "docx"
    assert body["document"]["title"] == "Положение об отделе закупок"
    by_text = {f["text"]: f for f in body["fragments"]}

    first = by_text["1. Планирование закупок."]  # Word auto-numbering rebuilt
    assert first["clause"] == "1"
    assert first["kind"] == "list_item"
    assert first["section"] == "Положение об отделе закупок › Функции"
    assert first["ref"].startswith("п. 1, абзац ")
    assert by_text["2. Проведение тендеров."]["clause"] == "2"
    assert by_text["3.5. Контроль исполнения договоров."]["clause"] == "3.5"

    row = by_text["Отдел закупок | Закупки"]
    assert row["kind"] == "table_row"
    assert row["location"] == {"table": 1, "row": 2, "cells": ["Отдел закупок", "Закупки"]}
    assert row["ref"] == "таблица 1, строка 2"


def test_xlsx_rows_with_cell_ranges():
    body = _post("structure.xlsx", _xlsx()).json()
    assert body["document"]["format"] == "xlsx"
    assert body["document"]["sheets"] == [
        {"name": "Структура", "hidden": False, "rows": 2, "header": ["Код", "Подразделение", "Подчинение"]}
    ]
    row = body["fragments"][1]
    assert row["text"] == "101 | Отдел закупок | Финансовый департамент"
    assert row["location"]["row"] == 3  # empty row 2 skipped, numbering kept
    assert row["location"]["range"] == "A3:C3"
    assert row["ref"] == "лист «Структура», строка 3"


def test_pdf_merges_wrapped_lines_and_splits_clauses():
    pdf = _pdf(["1. Planning of", "procurement.", "2. Tenders."])
    body = _post("order.pdf", pdf).json()
    assert body["document"]["format"] == "pdf"
    assert body["document"]["pages"] == 1
    first, second = body["fragments"]
    assert first["text"] == "1. Planning of procurement."
    assert first["location"] == {"page": 1, "line_start": 1, "line_end": 2}
    assert first["ref"] == "п. 1, стр. 1, строки 1–2"
    assert second["clause"] == "2"


def test_pdf_without_text_layer_warns():
    body = _post("scan.pdf", _pdf([])).json()
    assert body["fragments"] == []
    assert any("OCR" in warning for warning in body["warnings"])


def test_format_detected_from_bytes_not_extension():
    body = _post("misnamed.pdf", _xlsx()).json()
    assert body["document"]["format"] == "xlsx"


@pytest.mark.parametrize(
    ("name", "data", "status", "code"),
    [
        ("notes.txt", b"plain text", 415, "unsupported_format"),
        ("legacy.doc", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 100, 415, "unsupported_format"),
        ("empty.docx", b"", 400, "empty_file"),
        ("broken.pdf", b"%PDF-1.4 garbage", 422, "unreadable_document"),
        ("broken.docx", b"PK\x03\x04garbage", 422, "unreadable_document"),
    ],
)
def test_bad_input(name, data, status, code):
    response = _post(name, data)
    assert response.status_code == status
    assert response.json()["error"]["code"] == code


def test_missing_file_field():
    response = client.post("/extract", data={"x": "1"})
    assert response.status_code == 400
    assert response.json()["error"]["code"] == "bad_request"


def test_zip_bomb_rejected(monkeypatch):
    from app import formats

    monkeypatch.setattr(formats, "MAX_UNZIPPED_BYTES", 1000)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "0" * 5000)
    response = _post("bomb.docx", buffer.getvalue())
    assert response.status_code == 413


def test_subitems_carry_marker_not_clause():
    """Sub-items "а)" must not open a clause (docs/TASK.md §4 clauses.js); they carry `marker`."""
    from docx.oxml.ns import qn

    document = docx.Document()
    document.add_paragraph("5.3.2. Директор департамента:")
    # Turn the template's "List Number" list into Russian letters: "а)", "б)".
    numbering = document.part.numbering_part.element
    style_num_id = document.styles["List Number"].element.pPr.numPr.numId.val
    num = next(n for n in numbering.findall(qn("w:num")) if n.get(qn("w:numId")) == str(style_num_id))
    abstract_id = num.find(qn("w:abstractNumId")).get(qn("w:val"))
    abstract = next(a for a in numbering.findall(qn("w:abstractNum")) if a.get(qn("w:abstractNumId")) == abstract_id)
    level = abstract.find(qn("w:lvl"))
    level.find(qn("w:numFmt")).set(qn("w:val"), "russianLower")
    level.find(qn("w:lvlText")).set(qn("w:val"), "%1)")
    document.add_paragraph("планирует аудит;", style="List Number")
    document.add_paragraph("контролирует качество.", style="List Number")
    document.add_paragraph("в) написано вручную.")
    buffer = io.BytesIO()
    document.save(buffer)

    fragments = _post("subitems.docx", buffer.getvalue()).json()["fragments"]
    assert [(f["text"], f["clause"], f["marker"]) for f in fragments] == [
        ("5.3.2. Директор департамента:", "5.3.2", None),
        ("а) планирует аудит;", None, "а"),
        ("б) контролирует качество.", None, "б"),
        ("в) написано вручную.", None, "в"),
    ]
    assert fragments[1]["ref"] == "абзац 2"


def test_pdf_subitem_marker():
    fragments = _post("order.pdf", _pdf(["5.3. Director:", "a) plans audits;", "b) reports."])).json()["fragments"]
    assert [(f["clause"], f["marker"]) for f in fragments] == [("5.3", None), (None, "a"), (None, "b")]
