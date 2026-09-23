"""Source-reference regressions using real OOXML numbering in generated DOCX files."""

import io

import docx
import pytest
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from fastapi.testclient import TestClient

from app.main import app


def _numbering(document, num_id, *, levels=None, overrides=""):
    root = document.part.numbering_part.element
    if levels is not None:
        root.append(parse_xml(
            f'<w:abstractNum {nsdecls("w")} w:abstractNumId="100">{levels}</w:abstractNum>'
        ))
    root.append(parse_xml(
        f'<w:num {nsdecls("w")} w:numId="{num_id}">'
        f'<w:abstractNumId w:val="100"/>{overrides}</w:num>'
    ))


def _paragraph(document, text, num_id, level=0):
    p = document.add_paragraph(text)
    num_pr = p._p.get_or_add_pPr().get_or_add_numPr()
    num_pr.get_or_add_numId().val = num_id
    num_pr.get_or_add_ilvl().val = level


def _extract(document):
    buffer = io.BytesIO()
    document.save(buffer)
    response = TestClient(app).post("/extract", files={"file": ("numbered.docx", buffer.getvalue())})
    assert response.status_code == 200
    return response.json()["fragments"]


@pytest.mark.parametrize("start", [0, 5])
def test_start_override_keeps_references_and_other_lists_independent(start):
    document = docx.Document()
    _numbering(document, 100, levels='''
        <w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/>
        <w:lvlText w:val="%1."/></w:lvl>''', overrides=f'''
        <w:lvlOverride w:ilvl="0"><w:startOverride w:val="{start}"/></w:lvlOverride>''')
    _numbering(document, 101)  # same abstract definition, no override
    _paragraph(document, "Проверять закупки.", 100)
    _paragraph(document, "Готовить отчёт.", 100)
    _paragraph(document, "Независимый список.", 101)

    fragments = _extract(document)
    assert [f["clause"] for f in fragments] == [str(start), str(start + 1), "1"]
    assert fragments[0]["text"] == f"{start}. Проверять закупки."
    assert fragments[0]["ref"] == f"п. {start}, абзац 1"
    assert fragments[2]["text"] == "1. Независимый список."


@pytest.mark.parametrize(("start_override", "expected"), [("", "3"), ('<w:startOverride w:val="5"/>', "5")])
def test_level_override_replaces_template_and_start_override_takes_precedence(start_override, expected):
    document = docx.Document()
    _numbering(document, 100, levels='''
        <w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/>
        <w:lvlText w:val="%1."/></w:lvl>''', overrides=f'''
        <w:lvlOverride w:ilvl="0">{start_override}
          <w:lvl w:ilvl="0"><w:start w:val="3"/><w:numFmt w:val="decimal"/>
          <w:lvlText w:val="%1)"/></w:lvl>
        </w:lvlOverride>''')
    _paragraph(document, "Согласовывать план.", 100)

    fragment = _extract(document)[0]
    assert fragment["text"] == f"{expected}) Согласовывать план."
    assert fragment["clause"] == expected
    assert fragment["ref"] == f"п. {expected}, абзац 1"


def test_multilevel_override_reapplies_when_parent_advances():
    document = docx.Document()
    _numbering(document, 100, levels='''
        <w:lvl w:ilvl="0"><w:start w:val="3"/><w:numFmt w:val="decimal"/>
          <w:lvlText w:val="%1."/></w:lvl>
        <w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="decimal"/>
          <w:lvlText w:val="%1.%2."/></w:lvl>''', overrides='''
        <w:lvlOverride w:ilvl="1"><w:startOverride w:val="7"/></w:lvlOverride>''')
    for level in [0, 1, 1, 0, 1]:
        _paragraph(document, "Функция подразделения.", 100, level)

    fragments = _extract(document)
    assert [f["clause"] for f in fragments] == ["3", "3.7", "3.8", "4", "4.7"]
    assert fragments[-1]["ref"] == "п. 4.7, абзац 5"
