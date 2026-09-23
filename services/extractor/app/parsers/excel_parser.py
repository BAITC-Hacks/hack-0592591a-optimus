"""Excel (.xlsx/.xlsm via openpyxl, legacy .xls via xlrd): one fragment per non-empty row.

Org-structure appendices are often a table "код | подразделение | подчинение |
функция". Each row keeps its cells with column letters so a finding can cite
"лист «Структура», строка 7".
"""

import datetime as dt
import io

from openpyxl.utils import get_column_letter

from ..model import ExtractionError, Fragment, ParseResult, clause_from_text, clause_ref, clean_text

MAX_ROWS = 20_000  # across all sheets


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, (dt.datetime, dt.date)):
        value = value.isoformat()
    return clean_text(value).replace("\n", " ")


def _rows_openpyxl(data: bytes):
    import openpyxl

    try:
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ExtractionError(422, "unreadable_document", "Не удалось прочитать файл Excel.") from exc
    try:
        for sheet in workbook.worksheets:
            yield sheet.title, sheet.sheet_state != "visible", sheet.iter_rows(values_only=True)
    finally:
        workbook.close()


def _rows_xlrd(data: bytes):
    import xlrd

    try:
        book = xlrd.open_workbook(file_contents=data)
    except Exception as exc:
        raise ExtractionError(422, "unreadable_document", "Не удалось прочитать файл Excel (.xls).") from exc
    for sheet in book.sheets():
        rows = (
            tuple(
                xlrd.xldate_as_datetime(cell.value, book.datemode) if cell.ctype == xlrd.XL_CELL_DATE else cell.value
                for cell in sheet.row(r)
            )
            for r in range(sheet.nrows)
        )
        yield sheet.name, sheet.visibility != 0, rows


def parse_excel(data: bytes, fmt: str) -> ParseResult:
    result = ParseResult()
    sheets_meta = []
    total = 0
    source = _rows_xlrd(data) if fmt == "xls" else _rows_openpyxl(data)
    for sheet_name, hidden, rows in source:
        header = None
        count = 0
        for row_no, values in enumerate(rows, start=1):
            cells = [
                {"column": get_column_letter(col), "value": text}
                for col, raw in enumerate(values, start=1)
                if (text := _cell_text(raw))
            ]
            if not cells:
                continue
            if total >= MAX_ROWS:
                result.warnings.append(f"Обработаны первые {MAX_ROWS} непустых строк; остальные строки пропущены.")
                break
            total += 1
            count += 1
            if header is None:
                header = [c["value"] for c in cells]
            text = " | ".join(c["value"] for c in cells)
            clause = clause_from_text(cells[0]["value"])
            span = cells[0]["column"] + str(row_no)
            if len(cells) > 1:
                span += f":{cells[-1]['column']}{row_no}"
            result.fragments.append(Fragment(
                kind="sheet_row",
                text=text,
                clause=clause,
                section=sheet_name,
                location={"sheet": sheet_name, "row": row_no, "range": span, "cells": cells},
                ref=clause_ref(clause, f"лист «{sheet_name}», строка {row_no}"),
            ))
        sheets_meta.append({"name": sheet_name, "hidden": hidden, "rows": count, "header": header})
        if total >= MAX_ROWS:
            break
    result.meta = {"sheets": sheets_meta}
    return result
